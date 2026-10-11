# SPDX-License-Identifier: MIT
"""Physical XLSX/CSV writer for validated workbooks; no projection logic.

Each file is staged and published atomically. A CSV and its Note sidecar are
two filesystem artifacts: a failure after the first publish is reported as
partial with the exact committed artifacts and their ordinary undo records.
"""
from __future__ import annotations

import csv
from contextlib import suppress
from datetime import date, datetime
import hashlib
import os
from pathlib import Path
import tempfile

from spreadsheet_cells import active_cell_count, neutralized_rows
from tabular.model import CAPABILITIES, Cell, Sheet, TableError, Workbook

_NOTE_MARKER = "metnos.tabular.notes.v1"


def _generated_note(worksheet) -> bool:
    comment = worksheet["A1"].comment
    return bool(comment and comment.author == "Metnos" and comment.text == _NOTE_MARKER)


def _digest(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def _native(cell: Cell):
    if cell.type == "date":
        return date.fromisoformat(cell.value)
    if cell.type == "datetime":
        return datetime.fromisoformat(cell.value)
    return cell.value


def _format(cell: Cell, currency: str = "") -> str:
    if cell.type == "date":
        return "yyyy-mm-dd"
    if cell.type == "datetime":
        return "yyyy-mm-dd hh:mm:ss"
    if cell.type == "currency":
        return (f'"{currency}" ' if currency else "") + "#,##0.00"
    return "General"


def _append_xlsx_row(worksheet, cells, columns) -> None:
    # Explicitly construct strings: assigning '=...' alone makes openpyxl
    # infer a formula. Do not touch any pre-existing user cells.
    from openpyxl.cell import Cell as XlsxCell
    row = []
    for position, cell in enumerate(cells):
        native = _native(cell)
        item = XlsxCell(worksheet, value=native)
        if isinstance(native, str):
            item.data_type = "s"
        item.number_format = _format(cell, columns[position].currency)
        row.append(item)
    worksheet.append(row)


def _has_data(worksheet) -> bool:
    return any(cell.value is not None for row in worksheet.iter_rows() for cell in row)


def _stage_xlsx(path: Path, staged: Path, workbook: Workbook, mode: str) -> int:
    if mode == "create":
        from backends.files.tabular_xlsx import stage_xlsx_create
        return stage_xlsx_create(staged, workbook)
    import openpyxl
    existing = path.is_file() and mode != "create"
    physical = openpyxl.load_workbook(path) if existing else openpyxl.Workbook()
    if not existing:
        physical.remove(physical.active)
    written = 0
    row_starts = {}
    try:
        requested_names = {sheet.name for sheet in workbook.sheets}
        physical_names = {name.casefold(): name for name in physical.sheetnames}
        if any(sheet.name.casefold() in physical_names
               and physical_names[sheet.name.casefold()] != sheet.name for sheet in workbook.sheets):
            raise TableError("existing_sheet_case_collision")
        for sheet in workbook.sheets:
            if sheet.notes and sheet.name in physical and not _generated_note(physical[sheet.name]):
                raise TableError("notes_destination_is_user_sheet")
        if mode == "overwrite" and any(_generated_note(ws) and ws.title not in requested_names for ws in physical):
            raise TableError("existing_notes_need_explicit_replacement")
        # Prevalidate all append schemas before mutating even the in-memory book.
        if mode == "append":
            for sheet in workbook.sheets:
                if sheet.name in physical and _has_data(physical[sheet.name]):
                    worksheet = physical[sheet.name]
                    actual = [cell.value for cell in worksheet[1]]
                    if ((sheet.has_header and actual != [column.header for column in sheet.columns])
                            or (not sheet.has_header and worksheet.max_column != len(sheet.columns))):
                        raise TableError("append_schema_mismatch")
                    if physical[sheet.name].max_row + len(sheet.rows) > 1_048_576:
                        raise TableError("append_row_limit")
        for sheet in workbook.sheets:
            present = sheet.name in physical
            worksheet = physical[sheet.name] if present else physical.create_sheet(sheet.name)
            if mode == "overwrite" and present:
                # Keep unrelated user sheets, including their formulas.
                position = physical.sheetnames.index(sheet.name)
                physical.remove(worksheet)
                worksheet = physical.create_sheet(sheet.name, position)
            populated = mode == "append" and _has_data(worksheet)
            header_rows = int(sheet.has_header and bool(sheet.columns))
            row_starts[sheet.name] = worksheet.max_row + 1 if populated else 1 + header_rows
            if not populated and header_rows:
                _append_xlsx_row(worksheet, [Cell("text", col.header) for col in sheet.columns], sheet.columns)
                written += len(sheet.columns)
            for row in sheet.rows:
                _append_xlsx_row(worksheet, row, sheet.columns)
                written += len(row)
            if mode != "append" or not populated:
                worksheet.freeze_panes = "A2" if header_rows else None
                if header_rows and sheet.rows:
                    worksheet.auto_filter.ref = worksheet.dimensions
            if sheet.notes:
                from openpyxl.comments import Comment
                worksheet["A1"].comment = Comment(_NOTE_MARKER, "Metnos")
        physical.save(staged)
    finally:
        physical.close()
    # Read back typed content before publication, including hostile strings.
    verify = openpyxl.load_workbook(staged, data_only=False)
    try:
        for sheet in workbook.sheets:
            worksheet = verify[sheet.name]
            start = row_starts[sheet.name]
            for offset, expected in enumerate(sheet.rows):
                for position, cell in enumerate(expected, 1):
                    actual = worksheet.cell(start + offset, position)
                    native = _native(cell)
                    # openpyxl reopens dates as midnight datetimes.
                    if isinstance(native, date) and not isinstance(native, datetime):
                        native = datetime.combine(native, datetime.min.time())
                    if (actual.value != native
                            and not (native == "" and actual.value is None)
                            and not (cell.type == "blank" and actual.value is None)):
                        raise TableError("xlsx_readback_mismatch")
                    if cell.type == "boolean" and (actual.data_type != "b" or not isinstance(actual.value, bool)):
                        raise TableError("xlsx_type_readback_mismatch")
                    if cell.type in {"number", "currency"} and (actual.data_type != "n"
                            or isinstance(actual.value, bool) or not isinstance(actual.value, (int, float))):
                        raise TableError("xlsx_type_readback_mismatch")
                    if cell.type in {"date", "datetime"} and actual.data_type != "d":
                        raise TableError("xlsx_type_readback_mismatch")
                    if isinstance(cell.value, str) and cell.type not in {"date", "datetime"} and actual.data_type == "f":
                        raise TableError("xlsx_formula_from_data")
    finally:
        verify.close()
    return written


def _csv_rows(sheet: Sheet) -> list[list]:
    headers = [[column.header for column in sheet.columns]] if sheet.has_header and sheet.columns else []
    return headers + [
        ["" if cell.value is None else cell.value for cell in row] for row in sheet.rows]


def _stage_csv(path: Path, staged: Path, sheet: Sheet, mode: str) -> tuple[int, int]:
    rows = _csv_rows(sheet)
    retained_rows = 0
    if mode == "append" and path.is_file() and path.stat().st_size:
        with path.open(newline="", encoding="utf-8") as previous:
            reader = csv.reader(previous)
            if sheet.has_header:
                header = next(reader, [])
                expected_header = neutralized_rows([[column.header for column in sheet.columns]])[0]
                if header != [str(value) for value in expected_header]:
                    raise TableError("append_schema_mismatch")
                retained_rows = 1
            for row in reader:
                if len(row) != len(sheet.columns):
                    raise TableError("append_schema_mismatch")
                retained_rows += 1
        if retained_rows + len(sheet.rows) > 1_048_576:
            raise TableError("append_row_limit")
        # Copy bytes, preserving the user's existing cells exactly.
        import shutil
        shutil.copyfile(path, staged)
        with staged.open("rb") as previous:
            previous.seek(-1, os.SEEK_END)
            needs_newline = previous.read(1) not in {b"\n", b"\r"}
        with staged.open("a", newline="", encoding="utf-8") as stream:
            if needs_newline:
                stream.write("\r\n")
            appended_rows = rows[1:] if sheet.has_header and sheet.columns else rows
            csv.writer(stream, dialect="excel").writerows(neutralized_rows(appended_rows))
        rows = appended_rows
    else:
        with staged.open("w", newline="", encoding="utf-8") as stream:
            csv.writer(stream, dialect="excel").writerows(neutralized_rows(rows))
    # Verify the exact newly written text before publishing either artifact.
    expected = [[str(value) for value in row] for row in neutralized_rows(rows)]
    with staged.open(newline="", encoding="utf-8") as stream:
        reader = csv.reader(stream)
        for _ in range(retained_rows):
            if next(reader, None) is None:
                raise TableError("csv_readback_mismatch")
        for row in expected:
            if next(reader, None) != row:
                raise TableError("csv_readback_mismatch")
        if next(reader, None) is not None:
            raise TableError("csv_readback_mismatch")
    return sum(map(len, rows)), active_cell_count(rows)


def write(workbook: Workbook, destination: dict, mode: str) -> dict:
    # A deserialized/forged contract must be rejected before any physical I/O.
    workbook = Workbook.from_dict(workbook.to_dict())
    backend = destination.get("backend")
    if backend not in {"local_xlsx", "csv"} or mode not in CAPABILITIES[backend].modes:
        raise TableError("unsupported_backend_mode")
    # Import at the I/O boundary; the compiler remains independent of config.
    from backends.files.local import (_resolve_spreadsheet_output_path, _sanitize_filename,
                                     _spreadsheet_out_dir)
    from executor_helpers import backup_file_for_undo
    from path_alias import check_mutating_path_ambiguity
    raw_path = destination.get("path") or destination.get("spreadsheet_id")
    path = (_resolve_spreadsheet_output_path(raw_path) if raw_path else
            _spreadsheet_out_dir() / (_sanitize_filename(workbook.title) + (".csv" if backend == "csv" else ".xlsx")))
    expected_suffix = ".csv" if backend == "csv" else ".xlsx"
    if path.suffix.lower() != expected_suffix:
        raise TableError("destination_suffix_mismatch")
    if path.is_symlink():
        raise TableError("symlink_destination")
    artifacts = [(path, workbook)]
    if backend == "csv":
        data_sheets = [sheet for sheet in workbook.sheets if not sheet.notes]
        note_sheets = [sheet for sheet in workbook.sheets if sheet.notes]
        if len(data_sheets) != 1 or len(note_sheets) > 1:
            raise TableError("csv_multiple_data_sheets")
        artifacts = [(path, data_sheets[0])]
        if note_sheets:
            note_path = path.with_name(path.stem + "_note.csv")
            if destination.get("notes_path") and Path(destination["notes_path"]) != note_path:
                raise TableError("notes_path_mismatch")
            if note_path.exists() and mode != "create" and not destination.get("notes_path"):
                raise TableError("existing_notes_scope_required")
            artifacts.append((note_path, note_sheets[0]))
        elif mode != "create" and path.with_name(path.stem + "_note.csv").exists():
            raise TableError("existing_notes_need_explicit_replacement")
    # Validate every path before backup, mkdir or writing either CSV artifact.
    for artifact, _ in artifacts:
        ambiguity = check_mutating_path_ambiguity(artifact, target_must_exist=mode != "create")
        if ambiguity:
            raise TableError("ambiguous_destination")
        if artifact.is_symlink() or (artifact.exists() and not artifact.is_file()):
            raise TableError("invalid_destination")
        if mode == "create" and artifact.exists():
            raise TableError("destination_exists")
    if mode != "create" and not path.is_file():
        raise TableError("destination_not_found")
    staged_files = []
    results = []
    try:
        for artifact, table in artifacts:
            artifact.parent.mkdir(parents=True, exist_ok=True)
            existed = artifact.is_file()
            digest = _digest(artifact) if existed else None
            blob = backup_file_for_undo(artifact) if existed else None
            descriptor, temporary = tempfile.mkstemp(prefix=".metnos-table-", suffix=expected_suffix, dir=artifact.parent)
            os.close(descriptor)
            staged = Path(temporary)
            staged_files.append((artifact, staged, existed, digest, blob, table))
            if backend == "csv":
                count, prefixed = _stage_csv(artifact, staged, table, mode)
            else:
                count, prefixed = _stage_xlsx(artifact, staged, workbook, mode), 0
            # Persist complete staged bytes before the atomic link/replace.
            with staged.open("rb") as stream:
                os.fsync(stream.fileno())
            staged_files[-1] += (count, prefixed, _digest(staged))
        for artifact, staged, existed, digest, blob, table, count, prefixed, published_digest in staged_files:
            # Build the whole undo/result row before the mutation. Once link or
            # replace returns, record it before cleanup or another fallible I/O.
            row = {"ok": True, "path": str(artifact), "spreadsheet_id": str(artifact),
                   "created": not existed, "kind": "spreadsheet", "mode": mode,
                   "updated_cells": count, "sha256": published_digest,
                   "text_prefixed_cells": prefixed,
                   "notes": backend == "csv" and table.notes}
            row["data_rows"] = (sum(len(sheet.rows) for sheet in workbook.sheets if not sheet.notes)
                                if backend != "csv" else (0 if table.notes else len(table.rows)))
            if blob:
                row.update(prev_blob_path=blob, prev_blob_sha256=digest)
            if existed:
                if not artifact.is_file() or _digest(artifact) != digest:
                    raise TableError("destination_changed_during_write")
                os.replace(staged, artifact)
            else:
                # Atomic exclusive publication, including overwrite's new Note.
                os.link(staged, artifact)
            results.append(row)
            with suppress(OSError):
                staged.unlink(missing_ok=True)
        return {"ok": True, "results": results, "used": len(results), "verified": True,
                "spreadsheet_id": str(path), "path": str(path), "files_source": "local",
                "n_created": sum(row["created"] for row in results),
                "n_written": len(results), "_undo": _undo(results)}
    except Exception as exc:
        # Committed artifacts remain visible and undoable if a later publish
        # failed. Never hide them or return an all-or-nothing success fiction.
        if results:
            return {"ok": False, "partial": True, "partial_reason": "artifact_publish_failed",
                    "error_code": "ERR_TABLE_WRITE", "error_reason":
                        "destination_exists" if isinstance(exc, FileExistsError) else getattr(exc, "reason", "io_error"),
                    "results": results, "used": len(results), "verified": False,
                    "n_created": sum(row["created"] for row in results), "n_written": len(results),
                    "files_source": "local", "_undo": _undo(results)}
        if isinstance(exc, FileExistsError):
            raise TableError("destination_exists") from exc
        raise
    finally:
        for staged_item in staged_files:
            with suppress(OSError):
                staged_item[1].unlink(missing_ok=True)


def _undo(results: list[dict]) -> dict:
    patterns = []
    if any(row.get("prev_blob_path") for row in results):
        patterns.append("restore_blob_backup")
    if any(row.get("created") for row in results):
        patterns.append("delete_created_paths")
    return {"reverse_pattern": patterns, "scope": {"client": "local"},
            "paths": [row["path"] for row in results]}
