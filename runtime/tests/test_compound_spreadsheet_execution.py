"""Gate d'ESECUZIONE §2.8 per il compound «leggi/estrai X → costruisci un foglio».

Il banco `bench/compound_extract_create_bench.py` verifica la FORMA del piano
(read→extract→create). Qui verifichiamo l'ESECUZIONE deterministica del
CONTRATTO I/O executor — la parte che falliva in SILENZIO (§2.8): un piano
strutturalmente giusto che produce un foglio HEADER-ONLY (dati persi) o 0 record.

Casi (tutti DETERMINISTICI, niente LLM — sono i rami §7.9):
  1. create con entries=RIGHE (list[list], output di read_spreadsheet §2.6) NON
     scarta i dati → foglio popolato (FIX-4, turn 3da933e5-family).
  2. extract su RIGHE con HEADER → proiezione deterministica dei `fields`
     (structured_projection): righe-foglio → record coi campi richiesti.
  3. extract su RIGHE HEADERLESS → le celle diventano testo per il ramo LLM
     (nessun 0-record muto).

Radice non-ovvia (vedi project_compound_spreadsheet_reliability): la conversione
riga→record vive nel CONSUMER (extract_entries), NON in read_spreadsheet — cosi'
l'output di read resta list[list] per describe/filter/matcher (zero ripple).
"""
from __future__ import annotations

import sys
from pathlib import Path

_RUNTIME = str(Path(__file__).resolve().parent.parent)
if _RUNTIME not in sys.path:
    sys.path.insert(0, _RUNTIME)


# --- 1. create preserva le RIGHE (FIX-4): mai foglio header-only -------------

def test_create_spreadsheet_preserves_list_of_rows(tmp_path):
    from backends.files import local

    out_path = tmp_path / "out.xlsx"
    rows = [["data", "descrizione", "importo"],
            ["2026-07-01", "spesa", "10,5"],
            ["2026-07-02", "bar", "3,2"]]
    res = local.create_spreadsheet({
        "title": "spese", "path": str(out_path),
        "entries": rows, "columns": ["data", "descrizione", "importo"],
    })
    assert res["ok"] is True
    import openpyxl
    got = [list(r) for r in openpyxl.load_workbook(
        str(out_path), read_only=True, data_only=True).active.iter_rows(values_only=True)]
    # §2.8: NON header-only — le 3 righe (header + 2 dati) preservate.
    assert len(got) >= 3, f"foglio header-only (dati persi): {got}"
    assert got[1] == ["2026-07-01", "spesa", "10,5"]


# --- 2. extract su righe con HEADER → proiezione deterministica --------------

def test_extract_entries_rows_headed_projection():
    from extract_entries import handle_extract_entries

    rows = [["data", "descrizione", "importo"],
            ["2026-07-01", "spesa", "10,5"],
            ["2026-07-02", "bar", "3,2"]]
    out = handle_extract_entries({"entries": rows, "fields": ["data", "importo"]})
    assert out["ok"] is True
    # deterministico (niente LLM) e proiezione ai SOLI campi richiesti.
    assert out.get("source") == "structured_projection"
    assert out["meta"]["deterministic"] is True
    assert out["used"] == 2, f"0/parziali record da righe con header: {out}"
    assert out["entries"][0] == {"data": "2026-07-01", "importo": "10,5"}


# --- 3. extract su righe HEADERLESS → testo per il ramo LLM ------------------

def test_extract_entries_rows_headerless_yields_text():
    from extract_entries import _pick_text, _rows_to_records

    headerless = [["3/7", "spesa", "111,18"], ["2/7", "bar", "10,50"]]
    # niente header → resta list[list] (il ramo LLM unisce le celle).
    assert _rows_to_records(headerless) == headerless
    txt = _pick_text(["3/7", "spesa", "111,18"])
    assert "3/7" in txt and "spesa" in txt and "111,18" in txt


# --- 4. round-trip: read_spreadsheet → extract(proj) → create ----------------

def test_roundtrip_sheet_to_projected_sheet(tmp_path):
    from backends.files import local
    from extract_entries import handle_extract_entries

    src = tmp_path / "src.xlsx"
    local.create_spreadsheet({
        "title": "src", "path": str(src),
        "values": [["data", "descrizione", "importo"],
                   ["2026-07-01", "spesa", "10,5"],
                   ["2026-07-02", "bar", "3,2"]],
    })
    read = local.read_spreadsheet({"spreadsheet_id": str(src)})
    assert read["ok"] is True
    # read_spreadsheet emette RIGHE (list[list]) — contratto §2.6 invariato.
    assert isinstance(read["entries"][0], list)
    proj = handle_extract_entries({"entries": read["entries"],
                                   "fields": ["data", "importo"]})
    assert proj["used"] == 2
    out = tmp_path / "out.xlsx"
    cre = local.create_spreadsheet({"title": "out", "path": str(out),
                                    "entries": proj["entries"],
                                    "columns": ["data", "importo"]})
    assert cre["ok"] is True
    import openpyxl
    got = [list(r) for r in openpyxl.load_workbook(
        str(out), read_only=True, data_only=True).active.iter_rows(values_only=True)]
    assert len(got) >= 3, f"foglio proiettato header-only: {got}"
    assert got[0] == ["data", "importo"]
